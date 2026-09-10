import io
import re
from typing import Union
import docx
import pypdf

# Common bullet symbols regex
BULLET_PREFIX_REGEX = re.compile(r"^[\s\t]*([•\-*o▪■‣⁃–—\u2022\u25aa\u25a0\u2023\u2043\u2013\u2014\uf0a7\uf0b7\uf0a8]|\d+[\.\)])\s*")

def extract_text_from_pdf(file_input: Union[bytes, io.BytesIO]) -> str:
    """Extract text content from a PDF file byte stream into structured Markdown, preserving spacing gaps and bullet dots."""
    try:
        if isinstance(file_input, bytes):
            file_input = io.BytesIO(file_input)
        
        reader = pypdf.PdfReader(file_input)
        lines_out = []
        
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if not page_text:
                continue
                
            for raw_line in page_text.splitlines():
                line = raw_line.rstrip()
                stripped = line.strip()
                if not stripped:
                    lines_out.append("")
                    continue
                
                # Check for bullet points
                bullet_match = BULLET_PREFIX_REGEX.match(stripped)
                if bullet_match or stripped.startswith("- ") or stripped.startswith("* ") or stripped.startswith("• "):
                    clean_bullet = BULLET_PREFIX_REGEX.sub("", stripped).strip()
                    lines_out.append(f"• {clean_bullet}")
                else:
                    lines_out.append(line)
                    
        return "\n".join(lines_out).strip()
    except Exception as e:
        print(f"Error parsing PDF: {e}")
        return f"[Error extracting text from PDF: {str(e)}]"

def extract_text_from_docx(file_input: Union[bytes, io.BytesIO]) -> str:
    """Extract text content from a Word (.docx) file byte stream into structured Markdown, preserving spacing gaps and bullet dots."""
    try:
        if isinstance(file_input, bytes):
            file_input = io.BytesIO(file_input)
        
        doc = docx.Document(file_input)
        full_lines = []
        
        for para in doc.paragraphs:
            raw_text = para.text
            text = raw_text.strip()
            
            # Preserve empty line gaps and spacing between sections
            if not text:
                full_lines.append("")
                continue
                
            style_name = str(getattr(para.style, "name", "")).lower()
            
            # Reconstruct inline formatting (bold text) and preserve tab spaces for right-aligned dates
            runs_formatted = []
            for run in para.runs:
                r_text = run.text.replace("\t", "    ")
                if not r_text:
                    continue
                if run.bold and r_text.strip():
                    runs_formatted.append(f"**{r_text}**")
                else:
                    runs_formatted.append(r_text)
            
            line_content = "".join(runs_formatted).rstrip() if runs_formatted else text
            
            bullet_match = BULLET_PREFIX_REGEX.match(text)
            is_list_style = "list" in style_name or "bullet" in style_name
            
            # Structure detection
            if is_list_style or bullet_match:
                clean_bullet = BULLET_PREFIX_REGEX.sub("", line_content).strip()
                full_lines.append(f"• {clean_bullet}")
            else:
                full_lines.append(line_content)
                
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_lines.append(" | ".join(row_text))
                    
        return "\n".join(full_lines).strip()
    except Exception as e:
        print(f"Error parsing DOCX: {e}")
        return f"[Error extracting text from DOCX: {str(e)}]"

def parse_uploaded_file(file_obj) -> str:
    """
    Unified function to parse text from Streamlit UploadedFile objects
    or raw file bytes based on filename extension.
    """
    if file_obj is None:
        return ""
    
    filename = getattr(file_obj, "name", "").lower()
    
    try:
        # Read bytes from Streamlit UploadedFile or BytesIO
        if hasattr(file_obj, "getvalue"):
            file_bytes = file_obj.getvalue()
        elif hasattr(file_obj, "read"):
            file_bytes = file_obj.read()
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
        elif isinstance(file_obj, bytes):
            file_bytes = file_obj
        else:
            return str(file_obj)

        if not file_bytes:
            return ""

        if filename.endswith(".pdf"):
            return extract_text_from_pdf(file_bytes)
        elif filename.endswith(".docx") or filename.endswith(".doc"):
            return extract_text_from_docx(file_bytes)
        else:
            # Default to UTF-8 text parsing
            try:
                return file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1", errors="ignore")
    except Exception as e:
        return f"[Error reading uploaded file: {str(e)}]"
